# -*- coding: utf-8 -*-
r"""端到端：交付 xlsx 的那两列**真的**一个是句中切片、一个是词条形式吗？

**为什么非真跑不可**：`tests.py` 里原有那两条列断言，一条是**自指**的
（`hdr == list(DEFAULT_EXPORT.columns)`，拿配置跟自己比，改列照样绿），
一条**硬编码下标 `[2][6]`**（加列后读错列也可能照样绿）。两条都证明不了
「交付出去的那两列装的是对的东西」。

这条测试真跑 `run.py`（假 client，零 API 调用），读回产出的 xlsx，断言：

1. 表头恰好 9 列、名字与顺序逐字正确；
2. **第 4/5 列 100% 是各自整句的逐字子串** —— `final.json` 要拿它去原句里定位，
   定位不到就标不上。**这是 2026-09-19 改列的全部意义**，改之前约 45% 不是；
3. 第 8/9 列非空，且**至少有一行真的做了还原**（`col8 != col4`）——
   否则第 2 条就是空真：没有还原发生时，两列本来就一样。
   这一条是**防止本测试自己退化成橡皮图章**的，所以它红了要当回事。

顺带验：字段说明 markdown 写进了 out 目录、也进了 `results.zip`。

**零 API 调用**（假 client 不发 HTTP），`cache_root` 指向临时目录。
"""
import asyncio
import json
import os
import pathlib
import re
import shutil
import sys
import tempfile
import zipfile

import openpyxl

# 本脚本从 `e2e/` 子目录运行，先把仓库根加进 sys.path（原来它在根目录下）。
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

# ⚠ 本脚本**从不发真实请求**（假 client 或假 HTTP 层），但 `LLMClient.__init__`
#   会读一次 key 才肯构造。给个哑值即可 —— 环境变量优先于 key 文件，真 key 碰不到。
#   这样一来这个脚本**任何人拿到都能跑**，不需要先配一把真 key。
os.environ.setdefault("INFERERA_API_KEY", "sk-dummy-e2e-never-sends-a-request")
# 公开仓库里没有生产语料 —— 用 fixtures/ 里的**合成夹具**（内容全部虚构）。
from getterms import config as _cfg
_cfg.CORPUS_DIRS["es"] = (pathlib.Path(__file__).resolve().parent.parent
                          / "fixtures" / "corpus_es")


import getterms.run as R
from getterms.llm import CallResult, LLMClient

# ⚠ 本脚本直接 `import getterms.run` 再调 `R.main()`，**绕过了 `run.py` 底部
#   `if __name__ == "__main__": utf8_stdout()`**。于是 stdout 编码取决于你从哪个
#   控制台跑：UTF-8 下 exit 0，GBK 下会在 run.py 的 print 上抛 UnicodeEncodeError
#   而 exit 1 —— 同一个脚本两个结果，这是最坏的一种「验证」。
#   这里显式设一次，让成败只取决于代码。（2026-09-19 审计发现三个 e2e 都有这问题。）
from getterms.config import utf8_stdout

utf8_stdout()

FAILED = []
MADE = []

SPEC_COLUMNS = ["sent_id", "src_text", "tgt_text", "term_src", "term_tgt",
                "types", "note", "term_src_dict", "term_tgt_dict"]


def ck(name, ok, detail=""):
    print(f"  {'OK  ' if ok else 'FAIL'} {name}" + (f"   ({detail})" if detail else ""))
    if not ok:
        FAILED.append(name)


_CJK = re.compile(r"[一-鿿]+")

# 用**生产代码自己的**校验器来挑夹具：只有在「词典形 != 切片」**且能通过真校验**
# 时才算数。这样保证测试里确实发生还原 —— 否则「第 4/5 列可定位」那条是空真
# （没有还原时两列本来就相同）。
from getterms.extract import check_lemma_romance, load_lemmatizer  # noqa: E402

LEM = load_lemmatizer()


def _lemma(span: str) -> str:
    if LEM is None:
        return ""
    return " ".join(LEM.lemmatize(w, lang="es") for w in span.split())


def _pick_foreign(text: str) -> tuple[str, str]:
    """在外语句子里挑一个「能被合法还原」的切片，返回 (切片, 词典形)。

    找不到就退而求其次（词典形给同形，等于没还原）—— 那种情况下
    「至少一行真的做了还原」那条断言会红，**说明本测试失去了区分力**，
    要当回事而不是把断言删掉。
    """
    words = text.split()
    weak = None
    for n in (2, 1, 3):
        for i in range(len(words) - n + 1):
            span = " ".join(words[i:i + n]).strip(".,;:!?¡¿\"'()—-·")
            if len(span) < 4 or not span.strip():
                continue
            cand = _lemma(span)
            if not cand or cand == span:
                continue
            if check_lemma_romance(cand, span, "es", lemmatizer=LEM) is None:
                return span, cand
            weak = weak or (span, cand)
    if weak:
        return weak
    for w in words:
        s = w.strip(".,;:!?¡¿\"'()—-·")
        if len(s) >= 3:
            return s, s
    return "", ""


def _pick_zh(text: str) -> str:
    m = _CJK.search(text)
    return m.group(0)[:3] if m else ""


class CannedClient(LLMClient):
    """按真实句子返回 canned 术语；复用父类全部真实逻辑（锚定/词典形/导出）。

    ⚠ 词典形是**模型给的 `dict_form` 字段**，我们只做校验 —— 所以这里必须真的
    给出这个字段（缺了 `_fill_dict_form` 会记一条失败，还可能触发纠正重试）。
    第一版漏了它，于是永远没有还原、把测试自己变成了空真。
    """

    def __init__(self, *a, **kw):
        super().__init__(*a, **kw)
        self._memo = {}
        self._tmp_cache = tempfile.mkdtemp(prefix="deliv_cols_cache_")
        self.cache_root = pathlib.Path(self._tmp_cache)
        MADE.append(self._tmp_cache)

    async def call(self, messages, file_md5, batch_idx, batch_key=""):
        key = batch_idx % 10_000
        if batch_idx != key and key in self._memo:
            return self._memo[key]
        user = messages[-1]["content"]
        k = user.find('"sent_id"')
        if k < 0:
            return CallResult(ok=True, content="[]", finish_reason="stop",
                              attempts=1, model=self.price_key, effort=self.effort,
                              wire_fp=self.wire_fp)
        i, j = user.rfind("[", 0, k), user.rfind("]")
        pairs = json.loads(user[i:j + 1])
        items = []
        for p in pairs:
            src, tgt = str(p.get("src") or ""), str(p.get("tgt") or "")
            if _CJK.search(src):           # src=中文，外语在 tgt
                ts, tt = _pick_zh(src), ""
                fo_span, fo_dict = _pick_foreign(tgt)
                tt = fo_span
            else:                          # src=外语，中文在 tgt
                tt = _pick_zh(tgt)
                fo_span, fo_dict = _pick_foreign(src)
                ts = fo_span
            if not ts or not tt:
                continue
            items.append({"sent_id": p["sent_id"],
                          "terms": [{"term_src": ts, "term_tgt": tt,
                                     "dict_form": fo_dict,
                                     "types": ["loc"], "note": ""}]})
        r = CallResult(ok=True, content=json.dumps(items, ensure_ascii=False),
                       finish_reason="stop", attempts=1, model=self.price_key,
                       effort=self.effort, wire_fp=self.wire_fp)
        self._memo[key] = r
        return r


R.LLMClient = CannedClient

tmp = tempfile.mkdtemp(prefix="deliv_cols_out_")
argv = [
    "--lang", "es", "--model", "gemini-3.7-flash", "--effort", "high",
    "--limit-files", "1", "--limit-sentences", "20",
    "--out", tmp, "--tag", "delivcols",
    "--concurrency", "2", "--force", "--skip-preflight", "--no-correct-retry",
]
try:
    rc = R.main(argv)
    outd = pathlib.Path(tmp)
    xs = sorted(outd.glob("*term.xlsx"))
    print(f"  退出码 {rc}，产出 {len(xs)} 个 xlsx")

    if not xs:
        ck("有 xlsx 产出", False, "没有产出，后面无从验起")
    else:
        ws = openpyxl.load_workbook(xs[0]).active
        hdr = [c.value for c in ws[1]]
        ci = {n: i for i, n in enumerate(hdr) if n}
        print(f"  表头: {hdr}")
        ck("表头恰好 9 列、名字与顺序逐字正确", hdr == SPEC_COLUMNS, f"{hdr}")

        rows = [r for r in ws.iter_rows(min_row=2, values_only=True)]
        rows = [r for r in rows if r and r[ci["term_src"]]]
        ck("产出了行（断言不是空真）", len(rows) >= 5, f"{len(rows)} 行")

        bad = []
        for r in rows:
            for tc, xc in (("term_src", "src_text"), ("term_tgt", "tgt_text")):
                term = str(r[ci[tc]] or "")
                sent = str(r[ci[xc]] or "")
                if term and term not in sent:
                    bad.append(f"{tc}={term!r} ∉ {xc}")
        ck("**第 4/5 列 100% 是句中逐字子串**（改之前约 45% 不是）",
           not bad, f"{len(rows)} 行中 {len(bad)} 行定位不到"
                    + (f"，例：{bad[0]}" if bad else ""))

        empty = [r for r in rows
                 if not str(r[ci["term_src_dict"]] or "")
                 or not str(r[ci["term_tgt_dict"]] or "")]
        ck("第 8/9 列都非空", not empty, f"{len(empty)} 行有空列")

        # ⚠ 外语可能在任一侧（本文件是 src=西语、tgt=中文），所以要**两侧都看**。
        #   第一版只比了 `term_tgt`（中文侧），中文无屈折、那一列恒等于切片，
        #   于是永远算出「0 行还原」—— 断言写错了侧，不是流水线没还原。
        def _side_changed(r):
            ds, ss = str(r[ci["term_src_dict"]] or ""), str(r[ci["term_src"]] or "")
            dt, st = str(r[ci["term_tgt_dict"]] or ""), str(r[ci["term_tgt"]] or "")
            return (ds != ss) or (dt != st)

        changed = [r for r in rows if _side_changed(r)]
        _ex = ""
        for r in changed[:1]:
            for a, b in (("term_src", "term_src_dict"), ("term_tgt", "term_tgt_dict")):
                if str(r[ci[a]] or "") != str(r[ci[b]] or ""):
                    _ex = f"，例 {r[ci[a]]!r} -> {r[ci[b]]!r}"
                    break
        print(f"  真的做了还原的行: {len(changed)}/{len(rows)}{_ex}")
        ck("**至少一行真的做了还原**（否则上面那条是空真、本测试退化成橡皮图章）",
           len(changed) >= 1,
           "没有还原发生 -> 两列本来就相同，第 4/5 列那条断言不携带信息")

        # 把这份产出直接喂给交付门禁：证明 verify 在新格式下**真的在判**，
        # 而不是崩掉（列名一改就 KeyError）或变成橡皮图章（恒绿灯）。
        from getterms import verify as V
        idx, note = V.load_span_index(outd)
        probs, vstats = V.check_xlsx(xs[0], lang="es", span_idx=idx)
        print(f"  verify: {note}；词典形档 exact={vstats['exact']} "
              f"morph={vstats['morph']} neither={vstats['neither']}")
        ck("verify 在新格式上不报问题", not probs, "；".join(probs)[:200])
        ck("verify 的形态档**不是橡皮图章**：morph 真的数到了还原",
           vstats["morph"] >= 1,
           f"morph={vstats['morph']} —— 为 0 说明门禁恒绿灯（交付值恒等于切片时会这样）")

        # 字段说明：写进 out 目录 + 进 zip
        spec = outd / "术语表字段说明.md"
        ck("字段说明 markdown 写进了 out 目录", spec.exists(), str(spec.name))
        zp = outd / "results.zip"
        if zp.exists():
            with zipfile.ZipFile(zp) as z:
                names = z.namelist()
            ck("字段说明也进了 results.zip",
               "术语表字段说明.md" in names, f"zip 内 {len(names)} 个文件")
            ck("zip 里 xlsx 仍是扁平 basename（没有带目录）",
               all("/" not in n for n in names), f"{names[:3]}")
        else:
            ck("产出了 results.zip", False, "没找到 results.zip")
finally:
    shutil.rmtree(tmp, ignore_errors=True)
    for d in MADE:
        shutil.rmtree(d, ignore_errors=True)

print()
if FAILED:
    print(f"端到端失败 {len(FAILED)} 项：")
    for x in FAILED:
        print(f"  - {x}")
    sys.exit(1)
print("端到端通过：第 4/5 列可逐字定位、第 8/9 列是词条形式、说明随包走")
