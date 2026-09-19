# -*- coding: utf-8 -*-
r"""端到端：并发跑完之后，xlsx 的行序真的是句序吗？

**为什么非跑不可**：行序修复（`run.py` 的 `_sid_key`）此前只有**源码字符串断言**
（`tests.py:491` 检查 `"_rows.sort(key=_sid_key)" in src`）。那正是 429 熔断第一版
栽的坑 —— 断言全过、实测无效。字符串在源码里 ≠ 它会在运行时被执行、会真的改变输出。

这里真跑 `run.py`，只把 `LLMClient.call` 换成一个「**批内逆序返回**」的假实现，
并且让**第 0 批慢、第 1 批快**（复现旧 bug 的成因：批完成顺序 ≠ 句序）。
然后用真的导出路径产出 xlsx，读回来数 sent_id。

- 若排序生效：sent_id 单调不减（本文件是 1..20 连续）。
- 若不生效：先完成的第 1 批（11~20，且批内还逆序）整块排在第 0 批前面。

**零 API 调用**（假 client 不发 HTTP），也不碰真 `cache/`（cache_root 指向临时目录）。
"""
import asyncio
import json
import os
import pathlib
import shutil
import sys
import tempfile

import openpyxl

# 本脚本从 `e2e/` 子目录运行，先把仓库根加进 sys.path（原来它在根目录下）。
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

# ⚠ 本脚本**从不发真实请求**（假 client 或假 HTTP 层），但 `LLMClient.__init__`
#   会读一次 key 才肯构造。给个哑值即可 —— 环境变量优先于 key 文件，真 key 碰不到。
#   这样一来这个脚本**任何人拿到都能跑**，不需要先配一把真 key。
os.environ.setdefault("INFERERA_API_KEY", "sk-dummy-e2e-never-sends-a-request")
# 公开仓库里没有生产语料 —— 用 fixtures/ 里的**合成夹具**（内容全部虚构）。
from getterms import config as _cfg
_cfg.CORPUS_DIRS["es"] = (pathlib.Path(__file__).resolve().parent.parent
                          / "fixtures" / "corpus_es")


import getterms.run as R
from getterms.llm import CallResult, LLMClient

# ⚠ 绕过了 `run.py` 底部的 `if __name__ == "__main__": utf8_stdout()` ——
#   不显式设一次，本脚本的成败会取决于控制台编码（UTF-8 过 / GBK 挂）。
from getterms.config import utf8_stdout

utf8_stdout()

FAILED = []


def ck(name, ok, detail=""):
    print(f"  {'OK  ' if ok else 'FAIL'} {name}" + (f"   ({detail})" if detail else ""))
    if not ok:
        FAILED.append(name)


class ReversingClient(LLMClient):
    """按批号人为制造「完成顺序 ≠ 句序」+「批内逆序」。复用父类全部真实逻辑。"""

    def __init__(self, *a, **kw):
        super().__init__(*a, **kw)
        self._memo = {}
        self._tmp_cache = tempfile.mkdtemp(prefix="roworder_cache_")
        self.cache_root = pathlib.Path(self._tmp_cache)
        MADE.append(self._tmp_cache)

    async def call(self, messages, file_md5, batch_idx, batch_key=""):
        # 纠正重试用 10_000+ 命名空间，消息里**没有** pairs（末条是纠正模板）。
        # 本测试用 --no-correct-retry，这条路不该走到；真走到了就回放首轮答案。
        key = batch_idx % 10_000
        if batch_idx != key and key in self._memo:
            return self._memo[key]
        user = messages[-1]["content"]
        k = user.find('"sent_id"')
        if k < 0:
            return CallResult(ok=True, content="[]", finish_reason="stop",
                              attempts=1, model=self.price_key, effort=self.effort,
                              wire_fp=self.wire_fp)
        i, j = user.rfind("[", 0, k), user.rfind("]")
        pairs = json.loads(user[i:j + 1])
        # 第 0 批慢、第 1 批快 -> 完成顺序反转（旧 bug 的成因）
        await asyncio.sleep(0.6 if batch_idx == 0 else 0.0)
        items = []
        for p in reversed(pairs):              # ← 批内也逆序，两个方向一起试
            src = str(p.get("src") or "").strip()
            tgt = str(p.get("tgt") or "").strip()
            if not src or not tgt:
                continue
            items.append({
                "sent_id": p["sent_id"],
                "terms": [{"term_src": src[:8], "term_tgt": tgt[:4],
                           "types": ["loc"], "note": ""}],
            })
        r = CallResult(
            ok=True, content=json.dumps(items, ensure_ascii=False),
            finish_reason="stop", attempts=1, latency_s=0.0,
            model=self.price_key, effort=self.effort, wire_fp=self.wire_fp)
        self._memo[key] = r
        # 记下**返回的先后与批内顺序** —— 这就是「不排序时会得到的行序」，
        # 断言产出必须**不等于**它，否则本测试就是空真。
        RET_ORDER.append([it["sent_id"] for it in items])
        return r


MADE = []
RET_ORDER = []          # 假 client 每次返回的 sent_id 顺序（= 不排序时的行序）
R.LLMClient = ReversingClient

tmp = tempfile.mkdtemp(prefix="roworder_out_")
argv = [
    "--lang", "es", "--model", "gemini-3.7-flash", "--effort", "high",
    "--limit-files", "1", "--limit-sentences", "20",
    "--out", tmp, "--tag", "roworder",
    "--concurrency", "2", "--force", "--no-zip", "--skip-preflight",
    "--no-correct-retry",
]
try:
    rc = R.main(argv)
    xs = sorted(pathlib.Path(tmp).glob("*term.xlsx"))
    print(f"  退出码 {rc}，产出 {len(xs)} 个 xlsx")
    if not xs:
        ck("有 xlsx 产出", False, "没有产出，后面无从验起")
    else:
        ws = openpyxl.load_workbook(xs[0]).active
        ids = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
        hdr = [c.value for c in ws[1]]
        print(f"  文件 {xs[0].name}")
        print(f"  表头 {hdr}")
        print(f"  sent_id 列: {ids}")

        ck("确实产出了行（不是空表 -> 后面的断言不是空真）", len(ids) >= 10,
           f"{len(ids)} 行")
        ck("**行序是句序**（单调不减）",
           all(isinstance(a, int) and isinstance(b, int) and a <= b
               for a, b in zip(ids, ids[1:])),
           f"{ids[:6]} … {ids[-4:]}")
        ck("首行是第 1 句（不是先跑完的那一批）",
           ids and ids[0] == 1, f"首值 {ids[0] if ids else None}")
        ck("没有重复/丢句（1..N 各一次）",
           ids == list(range(1, len(ids) + 1)),
           f"共 {len(ids)} 行")

        # 反向证据（可证伪的那一条）：假 client 的返回顺序 = 不排序时会写出的行序。
        # 若排序没生效，ids 会**等于**它；断言两者必须不同。
        natural = [sid for chunk in RET_ORDER for sid in chunk]
        print(f"  假 client 的返回顺序（= 不排序时的行序）: {natural}")
        ck("排序**真的改变了**行序（产出 ≠ 自然追加顺序）",
           ids != natural, f"自然序 {natural[:6]}… vs 产出 {ids[:6]}…")
        ck("且自然序确实是乱的（否则这条断言也是空真）",
           natural != sorted(natural), f"{natural[:8]}")
finally:
    shutil.rmtree(tmp, ignore_errors=True)
    for d in MADE:
        shutil.rmtree(d, ignore_errors=True)

print()
if FAILED:
    print(f"端到端失败 {len(FAILED)} 项：")
    for x in FAILED:
        print(f"  - {x}")
    sys.exit(1)
print("端到端通过：并发完成顺序与批内顺序都被人为打乱，导出后仍是句序")
